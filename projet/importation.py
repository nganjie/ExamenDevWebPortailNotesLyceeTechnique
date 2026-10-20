"""importation.py : lecture et vérification des fichiers d'élèves (CSV, XLSX, XLS, ODS).

Les fonctions sont indépendantes de Flask pour pouvoir être testées facilement.
"""

import csv
import io
import re
import unicodedata
from datetime import date, datetime

EXTENSIONS = (".csv", ".xlsx", ".xls", ".ods")
LIGNES_MAX = 500
MOTIF_MATRICULE = re.compile(r"^LT[0-9]{2}-F[0-9]-[0-9]{4}$")

# Champs attendus et noms de colonnes reconnus (comparés sans accents, espaces ni majuscules)
CHAMPS = {
    "nom": ("nom", "noms", "nomdefamille", "name", "lastname"),
    "prenom": ("prenom", "prenoms", "firstname"),
    "date_naissance": ("datenaissance", "datedenaissance", "naissance", "nele", "neele", "date", "dob"),
    "matricule": ("matricule", "numero", "numeromatricule", "id"),
}
LIBELLES_CHAMPS = {
    "nom": "Nom", "prenom": "Prénom", "date_naissance": "Date de naissance", "matricule": "Matricule",
}


class PointVirgule(csv.excel):
    """Dialecte CSV par défaut : séparateur point-virgule (Excel en français)."""
    delimiter = ";"


class ErreurFichier(Exception):
    """Fichier illisible ou d'un format non accepté."""


def simplifier(texte):
    """« Date de naissance » -> « datedenaissance » (sans accents, espaces ni ponctuation)."""
    sans_accents = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", sans_accents.lower())


def texte_cellule(valeur):
    """Convertit une cellule de tableur en texte (les dates deviennent AAAA-MM-JJ)."""
    if valeur is None:
        return ""
    if isinstance(valeur, (datetime, date)):
        return valeur.strftime("%Y-%m-%d")
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return str(valeur).strip()


def lire_csv(contenu):
    """Lit un CSV en UTF-8 (avec ou sans BOM) ou Windows-1252, séparateur détecté."""
    try:
        texte = contenu.decode("utf-8-sig")
    except UnicodeDecodeError:
        texte = contenu.decode("cp1252")
    try:
        dialecte = csv.Sniffer().sniff(texte[:2048], delimiters=";,\t")
    except csv.Error:
        dialecte = PointVirgule
    return [[cellule.strip() for cellule in ligne] for ligne in csv.reader(io.StringIO(texte), dialecte)]


def lire_xlsx(contenu):
    """Lit la première feuille d'un classeur Excel récent (.xlsx)."""
    from openpyxl import load_workbook
    classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    feuille = classeur.worksheets[0]
    lignes = [[texte_cellule(v) for v in ligne] for ligne in feuille.iter_rows(values_only=True)]
    classeur.close()
    return lignes


def lire_xls(contenu):
    """Lit la première feuille d'un ancien classeur Excel (.xls)."""
    import xlrd
    classeur = xlrd.open_workbook(file_contents=contenu)
    feuille = classeur.sheet_by_index(0)
    lignes = []
    for numero in range(feuille.nrows):
        ligne = []
        for cellule in feuille.row(numero):
            if cellule.ctype == xlrd.XL_CELL_DATE:
                ligne.append(texte_cellule(xlrd.xldate_as_datetime(cellule.value, classeur.datemode)))
            else:
                ligne.append(texte_cellule(cellule.value))
        lignes.append(ligne)
    return lignes


def lire_ods(contenu):
    """Lit la première feuille d'un classeur OpenDocument (.ods, LibreOffice)."""
    from odf.namespaces import OFFICENS, TABLENS
    from odf.opendocument import load
    from odf.table import Table, TableCell, TableRow

    document = load(io.BytesIO(contenu))
    feuilles = document.spreadsheet.getElementsByType(Table)
    if not feuilles:
        return []
    lignes = []
    for ligne_odf in feuilles[0].getElementsByType(TableRow):
        repetitions_ligne = int(ligne_odf.getAttrNS(TABLENS, "number-rows-repeated") or 1)
        ligne = []
        for cellule in ligne_odf.getElementsByType(TableCell):
            repetitions = int(cellule.getAttrNS(TABLENS, "number-columns-repeated") or 1)
            date_cellule = cellule.getAttrNS(OFFICENS, "date-value")
            valeur = date_cellule[:10] if date_cellule else str(cellule).strip()
            ligne += [valeur] * min(repetitions, 50)  # cellules vides répétées : on limite
        if any(ligne) and repetitions_ligne >= 1:
            lignes.append(ligne)
    return lignes


def lire_fichier(nom_fichier, contenu):
    """Renvoie les lignes non vides d'un fichier (liste de listes de textes)."""
    extension = "." + nom_fichier.rsplit(".", 1)[-1].lower() if "." in nom_fichier else ""
    lecteurs = {".csv": lire_csv, ".xlsx": lire_xlsx, ".xls": lire_xls, ".ods": lire_ods}
    if extension not in lecteurs:
        raise ErreurFichier("Format non accepté : envoyez un fichier CSV, XLSX, XLS ou ODS.")
    try:
        lignes = lecteurs[extension](contenu)
    except ImportError as erreur:
        raise ErreurFichier(f"La bibliothèque nécessaire pour lire ce format n'est pas installée ({erreur.name}).")
    except Exception:
        raise ErreurFichier("Le fichier est illisible ou abîmé.")
    lignes = [ligne for ligne in lignes if any(cellule for cellule in ligne)]
    if len(lignes) < 2:
        raise ErreurFichier("Le fichier doit contenir une ligne d'en-tête puis au moins un élève.")
    if len(lignes) - 1 > LIGNES_MAX:
        raise ErreurFichier(f"Le fichier contient trop de lignes ({LIGNES_MAX} au maximum).")
    return lignes


def deviner_correspondance(entete):
    """Associe chaque champ attendu à l'indice d'une colonne de l'en-tête (ou None)."""
    simplifiees = [simplifier(colonne) for colonne in entete]
    correspondance = {}
    for champ, noms in CHAMPS.items():
        correspondance[champ] = next((i for i, colonne in enumerate(simplifiees) if colonne in noms), None)
    return correspondance


def normaliser_date(texte):
    """Accepte JJ/MM/AAAA, JJ-MM-AAAA, JJ.MM.AAAA ou AAAA-MM-JJ ; renvoie AAAA-MM-JJ ou None."""
    texte = (texte or "").strip()[:10]
    for format_date in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texte, format_date).date().isoformat()
        except ValueError:
            continue
    return None


def analyser_lignes(lignes, correspondance, matricules_existants, places, aujourdhui=None):
    """Vérifie chaque ligne (sans l'en-tête) et renvoie une liste de dictionnaires :

    {numero, nom, prenom, date_naissance, matricule, etat ("valide" | "doublon" | "erreur"), message}
    `places` est le nombre de places libres dans la classe de destination.
    """
    aujourdhui = aujourdhui or date.today()

    def cellule(ligne, champ):
        indice = correspondance.get(champ)
        return ligne[indice].strip() if indice is not None and indice < len(ligne) else ""

    resultats = []
    vus = set()
    valides = 0
    for numero, ligne in enumerate(lignes[1:], start=2):
        nom = " ".join(cellule(ligne, "nom").split()).upper()
        prenom = " ".join(cellule(ligne, "prenom").split())
        texte_date = cellule(ligne, "date_naissance")
        date_naissance = normaliser_date(texte_date)
        matricule = cellule(ligne, "matricule").upper()
        etat, message = "valide", "Valide"

        if not nom or not prenom:
            etat, message = "erreur", "Nom manquant" if not nom else "Prénom manquant"
        elif len(nom) > 40 or len(prenom) > 40:
            etat, message = "erreur", "Nom ou prénom trop long"
        elif date_naissance is None:
            etat, message = "erreur", "Date invalide" if texte_date else "Date manquante"
        elif not 10 <= (aujourdhui - date.fromisoformat(date_naissance)).days / 365.25 <= 30:
            etat, message = "erreur", "Âge hors de 10 à 30 ans"
        elif matricule and not MOTIF_MATRICULE.match(matricule):
            etat, message = "erreur", "Matricule mal formé"
        elif matricule and matricule in matricules_existants:
            etat, message = "doublon", "Matricule déjà utilisé"
        elif matricule and matricule in vus:
            etat, message = "doublon", "Matricule en double dans le fichier"
        elif valides >= places:
            etat, message = "erreur", "Classe pleine"

        if etat == "valide":
            valides += 1
            if matricule:
                vus.add(matricule)
        resultats.append({"numero": numero, "nom": nom, "prenom": prenom,
                          "date_naissance": date_naissance or texte_date, "matricule": matricule,
                          "etat": etat, "message": message})
    return resultats


def modele_csv():
    """Contenu du modèle CSV à remplir (UTF-8 avec BOM pour Excel)."""
    return "﻿nom;prenom;date_naissance;matricule\nMBARGA;Paul;14/03/2011;\nNGO BASSA;Aline;02/09/2010;\n"


def modele_xlsx():
    """Contenu binaire du modèle XLSX à remplir."""
    from openpyxl import Workbook
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Eleves"
    feuille.append(["nom", "prenom", "date_naissance", "matricule"])
    feuille.append(["MBARGA", "Paul", date(2011, 3, 14), None])
    feuille.append(["NGO BASSA", "Aline", date(2010, 9, 2), None])
    for colonne, largeur in zip("ABCD", (20, 20, 16, 16)):
        feuille.column_dimensions[colonne].width = largeur
    for ligne in feuille.iter_rows(min_row=2, min_col=3, max_col=3):
        for cellule_date in ligne:
            cellule_date.number_format = "DD/MM/YYYY"
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
